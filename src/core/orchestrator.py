"""
Orchestrator v2.0
Управляет всеми агентами, проверяет документы, сохраняет правки пользователя.
"""
import json
import os
from datetime import datetime
from typing import Dict, List, Any, Optional
from risk_scoring_engine import RiskScoringEngine

class Orchestrator:
    def __init__(self):
        self.risk_engine = RiskScoringEngine()
        self.history_file = "/workspace/conversation_history.json"
        self.rules_file = "/workspace/agent_rules.json"
        self.load_history()
        self.load_rules()
        
    def load_history(self):
        """Загружает историю переписки."""
        if os.path.exists(self.history_file):
            with open(self.history_file, 'r', encoding='utf-8') as f:
                self.history = json.load(f)
        else:
            self.history = []
            
    def save_history(self):
        """Сохраняет историю переписки."""
        with open(self.history_file, 'w', encoding='utf-8') as f:
            json.dump(self.history, f, ensure_ascii=False, indent=2)
            
    def load_rules(self):
        """Загружает правила для агентов."""
        if os.path.exists(self.rules_file):
            with open(self.rules_file, 'r', encoding='utf-8') as f:
                self.rules = json.load(f)
        else:
            self.rules = []
            
    def save_rules(self):
        """Сохраняет новые правила."""
        with open(self.rules_file, 'w', encoding='utf-8') as f:
            json.dump(self.rules, f, ensure_ascii=False, indent=2)

    def is_user_feedback(self, text: str) -> bool:
        """Определяет, является ли сообщение правкой к предыдущему ответу."""
        feedback_keywords = [
            'исправь', 'поправь', 'измени', 'добавь', 'убери', 
            'не так', 'ошибка', 'замечание', 'правка', 'редактируй',
            'перепиши', 'скорректируй', 'учти', 'внеси'
        ]
        text_lower = text.lower()
        
        # Если есть ключевые слова правки И история не пустая
        if self.history and any(keyword in text_lower for keyword in feedback_keywords):
            return True
        return False

    def process_message(self, user_text: str, user_id: int) -> Dict[str, Any]:
        """
        Обрабатывает сообщение пользователя.
        Если это правка - обновляет последний ответ.
        Если новый запрос - создает новый анализ.
        """
        timestamp = datetime.now().isoformat()
        
        # Проверяем, является ли сообщение правкой
        if self.is_user_feedback(user_text):
            # Это правка к предыдущему документу
            last_entry = self.history[-1] if self.history else None
            
            if last_entry and last_entry.get('type') == 'analysis':
                # Сохраняем правку в историю
                feedback_entry = {
                    'timestamp': timestamp,
                    'type': 'feedback',
                    'user_id': user_id,
                    'content': user_text,
                    'reference_to': last_entry.get('timestamp')
                }
                self.history.append(feedback_entry)
                
                # Создаем новое правило для агентов на основе правки
                new_rule = {
                    'created': timestamp,
                    'trigger': last_entry.get('doc_type', 'general'),
                    'instruction': f"При обработке подобных документов учитывать правку пользователя: {user_text}",
                    'source': 'user_feedback'
                }
                self.rules.append(new_rule)
                self.save_rules()
                
                # Генерируем обновленный ответ с учетом правки
                updated_analysis = self.risk_engine.analyze_document(
                    last_entry.get('original_text', '') + f"\n[ПРАВКА ПОЛЬЗОВАТЕЛЯ: {user_text}]",
                    last_entry.get('doc_type', 'document')
                )
                
                response = {
                    'status': 'updated',
                    'message': '✅ Правка принята и учтена. Документ обновлен.',
                    'analysis': updated_analysis,
                    'rules_added': 1
                }
            else:
                response = {
                    'status': 'error',
                    'message': '❌ Не найден предыдущий документ для редактирования. Пожалуйста, отправьте документ сначала.'
                }
        else:
            # Это новый запрос
            analysis = self.risk_engine.analyze_document(user_text, "document")
            
            # Сохраняем в историю
            entry = {
                'timestamp': timestamp,
                'type': 'analysis',
                'user_id': user_id,
                'original_text': user_text,  # Сохраняем полный текст для будущих правок
                'doc_type': self._detect_doc_type(user_text),
                'risk_score': analysis['risk_score'],
                'risk_level': analysis['risk_level'],
                'full_analysis': analysis  # Сохраняем полный анализ
            }
            self.history.append(entry)
            
            response = {
                'status': 'new',
                'message': f'📄 Анализ завершен. Риск: {analysis["risk_level"]} ({analysis["risk_score"]}/10)',
                'analysis': analysis
            }
        
        self.save_history()
        return response

    def _detect_doc_type(self, text: str) -> str:
        """Определяет тип документа по ключевым словам."""
        text_lower = text.lower()
        if 'фас' in text_lower or 'жалоб' in text_lower:
            return 'fas_complaint'
        elif 'замен' in text_lower or 'материал' in text_lower:
            return 'material_change'
        elif 'договор' in text_lower or 'контракт' in text_lower:
            return 'contract'
        elif 'акт' in text_lower:
            return 'act'
        else:
            return 'general'

    def generate_docx(self, analysis: Dict[str, Any]) -> bytes:
        """Генерирует .docx файл с проектом документа на основе контекста."""
        try:
            from docx import Document
            from docx.shared import Pt, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            
            doc = Document()
            
            # Определяем тип документа для заголовка
            doc_type = analysis.get('doc_type', 'general')
            timestamp = analysis.get('timestamp', datetime.now().isoformat())[:10]
            
            if doc_type == 'fas_complaint':
                title = 'ДОПОЛНЕНИЕ К ЖАЛОБЕ В ФАС'
            elif doc_type == 'material_change':
                title = 'ПИСЬМО О ЗАМЕНЕ МАТЕРИАЛА'
            elif doc_type == 'contract':
                title = 'ЗАКЛЮЧЕНИЕ ПО ДОГОВОРУ'
            else:
                title = 'ЮРИДИЧЕСКИЙ ДОКУМЕНТ'
            
            # Заголовок
            heading = doc.add_heading(title, 0)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Дата
            doc.add_paragraph(f'Дата: {timestamp}')
            doc.add_paragraph(f'От: ООО «СТРОЙСТАНДАРТ»')
            doc.add_paragraph('_' * 50)
            
            # Основной текст - берем из draft_document или генерируем на основе анализа
            draft_text = analysis.get('draft_document', '')
            
            if draft_text:
                for line in draft_text.strip().split('\n'):
                    if line.strip():
                        doc.add_paragraph(line)
            else:
                # Если draft не сформирован, создаем структуру на основе рекомендаций
                doc.add_heading('Анализ ситуации', level=1)
                for issue in analysis.get('detected_issues', []):
                    doc.add_paragraph(f'• {issue}', style='List Bullet')
                
                doc.add_heading('Рекомендации', level=1)
                for rec in analysis.get('recommendations', []):
                    doc.add_paragraph(f'✓ {rec}', style='List Bullet')
            
            # Футер
            doc.add_page_break()
            footer = doc.add_paragraph('Документ сгенерирован автоматически ИИ-юристом ООО "СТРОЙСТАНДАРТ"')
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer.style.font.size = Pt(8)
            footer.style.font.italic = True
            
            # Сохраняем в буфер
            from io import BytesIO
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            
            return buffer.getvalue()
            
        except ImportError as e:
            # Если python-docx не установлен, возвращаем текстовую версию
            report_text = f"""
{title}
{'=' * len(title)}
Дата: {analysis.get('timestamp', '')[:10]}
Риск: {analysis.get('risk_level', 'N/A')} ({analysis.get('risk_score', 0)}/10)

ПРОЕКТ ДОКУМЕНТА:
{analysis.get('draft_document', 'Не сгенерирован')}

---
Сгенерировано ИИ-юристом ООО "СТРОЙСТАНДАРТ"
"""
            return report_text.encode('utf-8')

# Тест
if __name__ == "__main__":
    orch = Orchestrator()
    result = orch.process_message("Жалоба в ФАС на незаконные действия заказчика", 12345)
    print(f"Status: {result['status']}")
    print(f"Message: {result['message']}")
